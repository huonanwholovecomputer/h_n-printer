# -*- coding: utf-8 -*-
"""验收 · Office 转换沙箱三项加固（2026-10-25 真机验证后修，S2-03 批 A）

对应 docs/S2-03沙箱三引擎验证_2026-10-25.md §7.2 的第 1/4/5 项（按严重度重排为 1/2/3）：

  ① **env 套娃**：`HN_OFFICE_SANDBOX=1` 时 worker 继承该 env ⇒ 再判"需要沙箱" ⇒ 再起
     worker。实测 11.5 s 内 79 个并发 convert_worker、最终超时失败。
     修法：父进程给 worker 塞 `HN_OFFICE_SANDBOX_CHILD=1`，`should_sandbox()` 见到即 False。
  ② **超时被 180 s 卡死**：原来恒定 180 s，盖掉了 LibreOffice 原本 600 s 的上限 ⇒
     大文件"关沙箱能转完、开沙箱拒转"。修法：`min(max(180, MB*5), 600)`，env 仍可覆盖。
  ③ **worker 日志被丢弃**：成功路径丢掉 worker 的 stderr ⇒ 开沙箱后看不到"引擎选择"。
     修法：worker 侧把日志级别开到 INFO 走 stderr；父进程把它转进主日志。

判据都是"能红的"：把任一修复删掉，下面的行为/静态断言必须转红（见 --reverse 说明：
反向验证用 git HEAD 版本做基线，本文件自身不写死版本号）。

用法: PYTHONPATH= PYTHONUTF8=1 <python311> tests/verify_office_sandbox_hardening.py
exit 0 = 全绿
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
import time

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace",
                              line_buffering=True)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOOL = os.path.join(REPO, "local_print_tool")
PY = sys.executable
HERE = os.path.dirname(os.path.abspath(__file__))

ok_all = True


def check(name, cond, detail=""):
    global ok_all
    tag = "PASS" if cond else "FAIL"
    if not cond:
        ok_all = False
    line = f"[{tag}] {name}"
    if detail:
        line += f"  -- {detail}"
    print(line)
    return bool(cond)


def section(title):
    print(f"\n=== {title} ===")


# ────────────────────────────────────────────────────────────────
# ① env 套娃闸门（静态 + 行为）
# ────────────────────────────────────────────────────────────────
section("① env 套娃：worker 内 should_sandbox 必须恒 False")

sys.path.insert(0, TOOL)
os.environ["HN_OFFICE_SANDBOX"] = "1"          # 强制开（env 通路，优先级最高）
os.environ.pop("HN_OFFICE_SANDBOX_CHILD", None)

import converter  # noqa: E402

fake_docx = os.path.join(tempfile.gettempdir(), "_s2_03_fake.docx")

check("不带 CHILD 标记时本地 .docx 判定为需要沙箱（基线）",
      converter.should_sandbox(fake_docx) is True)

os.environ["HN_OFFICE_SANDBOX_CHILD"] = "1"
check("★ 带 CHILD 标记（已在沙箱 worker 内）时恒 False —— 否则会无限套娃",
      converter.should_sandbox(fake_docx) is False,
      "这是套娃的唯一闸门")
check("CHILD 标记下云端来源同样 False（不应因标记而放宽既有判定）",
      converter.should_sandbox(fake_docx, from_cloud=True) is False)
os.environ.pop("HN_OFFICE_SANDBOX_CHILD", None)

check("没有 CHILD 标记时云端来源仍 False（S1-04 强制 LibreOffice 口径不变）",
      converter.should_sandbox(fake_docx, from_cloud=True) is False)
os.environ["HN_OFFICE_SANDBOX"] = "0"
check("开关关闭（HN_OFFICE_SANDBOX=0）时仍 False（零回归）",
      converter.should_sandbox(fake_docx) is False)
os.environ["HN_OFFICE_SANDBOX"] = "1"
os.environ["HN_OFFICE_SANDBOX"] = "1"

import office_sandbox  # noqa: E402

# 缺实现时**断言失败**而不是抛 AttributeError（反向验证的红必须落在断言上）
_child_env = getattr(office_sandbox, "_child_env", None)
CHILD_ENV = getattr(office_sandbox, "CHILD_ENV", None)
check("★ 实现了『给 worker 塞子进程标记』的 env 构造（_child_env + CHILD_ENV）",
      callable(_child_env) and bool(CHILD_ENV),
      "修复前：无此函数 ⇒ 套娃闸门不存在")
if callable(_child_env) and CHILD_ENV:
    env_child = _child_env()
    check("★ 父进程给 worker 的环境里带 CHILD_ENV=1",
          env_child.get(CHILD_ENV) == "1", f"{CHILD_ENV}={env_child.get(CHILD_ENV)}")
    check("子进程环境继承原环境（PATH 仍在，不能把 worker 环境清空）",
          bool(env_child.get("PATH") or env_child.get("Path")))
_is_child = getattr(office_sandbox, "_is_child_process", None)
if CHILD_ENV:
    os.environ[CHILD_ENV] = "true"
child_ok = bool(callable(_is_child) and _is_child()) if CHILD_ENV else False
check("_is_child_process() 认 1/true（与 should_sandbox 同口径）", child_ok)
if CHILD_ENV:
    os.environ.pop(CHILD_ENV, None)

# ────────────────────────────────────────────────────────────────
# ② 超时按文件大小放大
# ────────────────────────────────────────────────────────────────
section("② 超时：min(max(180, MB*5), 600)，env 仍可覆盖")

_tf = getattr(office_sandbox, "_timeout_for_size", None)
_ts = getattr(office_sandbox, "_timeout_sec", None)
check("★ 实现了按文件大小取超时（_timeout_for_size / _timeout_sec）",
      callable(_tf) and callable(_ts), "修复前：恒定 180 s，无此实现")

if callable(_tf):
    for mb, want in [(0.5, 180.0), (10.0, 180.0), (36.0, 180.0), (40.0, 200.0),
                     (120.0, 600.0), (1000.0, 600.0)]:
        got = _tf(mb)
        check(f"{mb} MB → {want:.0f}s", abs(got - want) < 1e-6, f"实际 {got:.0f}s")
    check("下限不比改前更短（180 s）—— 小文件不会因换口径反而超时",
          float(getattr(office_sandbox, "MIN_TIMEOUT_SEC", 0)) >= 180.0)
    check("上限与 LibreOffice 原口径一致（600 s）",
          abs(float(getattr(office_sandbox, "MAX_TIMEOUT_SEC", 0)) - 600.0) < 1e-6)

def safe(fn):
    """调用可能不存在/签名不同的实现：异常一律当 None（红要落在断言上，不是 traceback）。"""
    try:
        return fn()
    except Exception:  # noqa: BLE001
        return None


if callable(_ts):
    os.environ.pop("HN_OFFICE_SANDBOX_TIMEOUT", None)
    check("默认（无 size）仍是 180 s（旧调用方零回归）",
          abs((safe(lambda: _ts()) or 0) - 180.0) < 1e-6)
    got = safe(lambda: _ts(40.0))
    check("按大小生效：40 MB → 200 s",
          got is not None and abs(got - 200.0) < 1e-6, f"实际 {got}")
    os.environ["HN_OFFICE_SANDBOX_TIMEOUT"] = "30"
    got = safe(lambda: _ts(40.0))
    check("★ env 覆盖优先（HN_OFFICE_SANDBOX_TIMEOUT=30 → 30 s）",
          got is not None and abs(got - 30.0) < 1e-6, f"实际 {got}")
    os.environ["HN_OFFICE_SANDBOX_TIMEOUT"] = "abc"
    got = safe(lambda: _ts(40.0))
    check("畸形 env 不变成'没有超时'（回落到按大小）",
          got is not None and abs(got - 200.0) < 1e-6, f"实际 {got}")
    os.environ.pop("HN_OFFICE_SANDBOX_TIMEOUT", None)

# ────────────────────────────────────────────────────────────────
# ③ 真转换：沙箱开 + 日志转出（行为）
# ────────────────────────────────────────────────────────────────
section("③ 真转换：沙箱开时能转完、只起 1 个 worker、引擎日志能转进主日志")

# 静态判据（反向验证时即使跳过真转换，修复③也必须能红）
_sand_src = open(os.path.join(TOOL, "office_sandbox.py"), encoding="utf-8").read()
_worker_src = open(os.path.join(TOOL, "convert_worker.py"), encoding="utf-8").read()
check("★ 父进程有把 worker stderr 转进日志的实现与调用点",
      _sand_src.count("def _log_worker_stderr") == 1
      and _sand_src.count("_log_worker_stderr(") >= 2,
      "修复前：成功路径直接丢弃 stderr")
check("★ worker 侧把日志开到 INFO 并走 stderr",
      "basicConfig(stream=sys.stderr" in _worker_src
      and "level=logging.INFO" in _worker_src,
      "修复前：INFO 无 handler ⇒ 引擎选择不可见")

if os.environ.get("HN_SKIP_REAL_CONVERT") == "1":
    print("[NOTE] HN_SKIP_REAL_CONVERT=1：跳过真实转换（反向验证用；静态判据仍有效）")
    print("\n" + "=" * 56)
    print("Office 沙箱三项加固验收:", "全部通过 ✅" if ok_all else "存在失败 ❌")
    sys.exit(0 if ok_all else 1)

tmp = tempfile.mkdtemp(prefix="s2_03_hard_")
src_docx = os.path.join(tmp, "sample.docx")
out_pdf = os.path.join(tmp, "out.pdf")
runner = os.path.join(tmp, "runner.py")
log_path = os.path.join(tmp, "parent.log")

try:
    from docx import Document
    doc = Document()
    doc.add_heading("S2-03 批 A 沙箱加固验收", level=1)
    for i in range(20):
        doc.add_paragraph(f"第 {i + 1} 行：沙箱加固后的真实转换（中文 / ASCII mix 12345）")
    doc.save(src_docx)

    with open(runner, "w", encoding="utf-8") as f:
        f.write('''# -*- coding: utf-8 -*-
import logging, os, sys
sys.path.insert(0, r"%s")
logging.basicConfig(stream=sys.stderr, level=logging.INFO,
                    format="%%(levelname)s %%(name)s: %%(message)s")
import converter
conv = converter.UniversalConverter()
out = conv.convert(r"%s", r"%s")
print("CONVERT_OK", out, os.path.getsize(out))
''' % (TOOL, src_docx, out_pdf))

    env = dict(os.environ)
    env["HN_OFFICE_SANDBOX"] = "1"
    env.pop("HN_OFFICE_SANDBOX_CHILD", None)
    env["PYTHONPATH"] = ""
    env["PYTHONIOENCODING"] = "utf-8"

    ps_count = (
        'powershell.exe -NoProfile -Command '
        '"@(Get-CimInstance Win32_Process | Where-Object { '
        "$_.Name -like 'python*' -and $_.CommandLine -like '*convert_worker.py*' }).Count\""
    )

    def worker_count():
        try:
            r = subprocess.run(ps_count, shell=True, capture_output=True,
                               text=True, timeout=25)
            return int((r.stdout or "0").strip() or 0)
        except Exception:
            return None

    t0 = time.time()
    proc = subprocess.Popen([PY, runner], env=env, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE)
    peak = 0
    sampler_ok = True
    while proc.poll() is None and time.time() - t0 < 180:
        n = worker_count()
        if n is None:
            sampler_ok = False
        else:
            peak = max(peak, n)
        time.sleep(0.4)
    try:
        out, err = proc.communicate(timeout=30)
    except subprocess.TimeoutExpired:
        proc.kill()
        out, err = proc.communicate()
    dur = time.time() - t0
    stdout = (out or b"").decode("utf-8", "replace")
    stderr = (err or b"").decode("utf-8", "replace")
    with open(log_path, "w", encoding="utf-8") as f:
        f.write(stdout + "\n--- stderr ---\n" + stderr)

    check("沙箱开时真转换成功（< 180 s）",
          "CONVERT_OK" in stdout, f"rc={proc.returncode} {dur:.1f}s")
    check("产物 PDF 存在且非空",
          os.path.isfile(out_pdf) and os.path.getsize(out_pdf) > 0,
          f"{os.path.getsize(out_pdf) if os.path.isfile(out_pdf) else 0} B")
    if sampler_ok:
        check("★ 全程 convert_worker 进程数 ≤ 1（套娃的现场判据：实测 79）",
              peak <= 1, f"峰值 {peak}")
    else:
        print("[NOTE] 进程采样器不可用，跳过峰值断言（静态闸门与转换成功仍有效）")

    check("★ worker 的 stderr 被转进主进程日志（原来成功路径直接丢弃）",
          "转换沙箱 worker stderr" in stderr,
          "日志里应出现『转换沙箱 worker stderr（N 行）』")
    check("★ 日志里能看到 worker 侧的引擎/转换信息（引擎选择可见）",
          ("引擎选择" in stderr) or ("worker 开始转换" in stderr),
          "期望『引擎选择: …』或『worker 开始转换: …』")
    check("超时按大小取值的日志行也在（可追溯到规则本身）",
          "转换沙箱超时按文件大小取" in stderr)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

print("\n" + "=" * 56)
print("Office 沙箱三项加固验收:", "全部通过 ✅" if ok_all else "存在失败 ❌")
sys.exit(0 if ok_all else 1)
